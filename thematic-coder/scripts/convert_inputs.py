#!/usr/bin/env python3
"""Convert source documents in inputs/ to markdown in inputs_md/."""

import os
import sys
from pathlib import Path

INPUTS_DIR = Path("inputs")
OUTPUTS_DIR = Path("inputs_md")


def convert_txt(path: Path) -> str:
    return path.read_text(encoding="utf-8", errors="replace")


def convert_docx(path: Path) -> str:
    from docx import Document
    doc = Document(path)
    lines = []
    for para in doc.paragraphs:
        style = para.style.name if para.style else ""
        text = para.text.strip()
        if not text:
            lines.append("")
            continue
        if style.startswith("Heading 1"):
            lines.append(f"# {text}")
        elif style.startswith("Heading 2"):
            lines.append(f"## {text}")
        elif style.startswith("Heading 3"):
            lines.append(f"### {text}")
        else:
            lines.append(text)
    return "\n\n".join(line for line in lines)


def convert_pdf(path: Path) -> str:
    from pypdf import PdfReader
    reader = PdfReader(str(path))
    pages = []
    for i, page in enumerate(reader.pages, 1):
        text = page.extract_text() or ""
        if text.strip():
            pages.append(f"<!-- Page {i} -->\n{text.strip()}")
    return "\n\n".join(pages)


def convert_pptx(path: Path) -> str:
    from pptx import Presentation
    prs = Presentation(path)
    slides = []
    for i, slide in enumerate(prs.slides, 1):
        texts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    t = para.text.strip()
                    if t:
                        texts.append(t)
        if texts:
            slides.append(f"## Slide {i}\n\n" + "\n\n".join(texts))
    return "\n\n".join(slides)


def convert_xlsx(path: Path) -> str:
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True)
    sheets = []
    for sheet in wb.worksheets:
        rows = []
        for row in sheet.iter_rows(values_only=True):
            cells = [str(c) if c is not None else "" for c in row]
            if any(c.strip() for c in cells):
                rows.append(" | ".join(cells))
        if rows:
            sheets.append(f"## Sheet: {sheet.title}\n\n" + "\n".join(rows))
    return "\n\n".join(sheets)


CONVERTERS = {
    ".txt": convert_txt,
    ".md": convert_txt,
    ".docx": convert_docx,
    ".pdf": convert_pdf,
    ".pptx": convert_pptx,
    ".xlsx": convert_xlsx,
}


def main():
    OUTPUTS_DIR.mkdir(exist_ok=True)

    if not INPUTS_DIR.exists():
        print("inputs/ directory not found.")
        return

    source_files = [
        f for f in INPUTS_DIR.iterdir()
        if f.is_file() and f.suffix.lower() in CONVERTERS
    ]

    if not source_files:
        print("No source files found in inputs/")
        return

    converted = 0
    skipped = 0
    failed = 0

    for source in sorted(source_files):
        out_path = OUTPUTS_DIR / (source.stem + ".md")
        if out_path.exists():
            print(f"  skip  {source.name} (already converted)")
            skipped += 1
            continue
        converter = CONVERTERS[source.suffix.lower()]
        try:
            text = converter(source)
            out_path.write_text(text, encoding="utf-8")
            print(f"  ok    {source.name} -> {out_path.name}")
            converted += 1
        except ImportError as e:
            print(f"  fail  {source.name}: missing dependency -- {e}")
            print(f"        run: pip install {e.name}")
            failed += 1
        except Exception as e:
            print(f"  fail  {source.name}: {e}")
            failed += 1

    print(f"\nDone. Converted: {converted}, Skipped: {skipped}, Failed: {failed}")


if __name__ == "__main__":
    main()
